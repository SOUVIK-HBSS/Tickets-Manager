import os
import json
import csv
import re
import threading
import time
from pathlib import Path
from flask import Flask, render_template, request, jsonify, send_file
from werkzeug.utils import secure_filename
try:
    import urllib.request
    import urllib.error
except ImportError:
    import urllib as urllib_request

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 200 * 1024 * 1024
app.config["SEND_FILE_MAX_AGE_DEFAULT"] = 0

DB_FILE = "ticket_db.json"
PPT_PLAIN = "ppt_output.txt"
PPT_HTML = "ppt_output.html"
PPT_HTML_ALL = "ppt_output_all.html"
GSHEET_LINKS_FILE = "gsheet_links.json"
GSHEET_CREDS_FILE = "gsheet_credentials.json"

STATUS_COLORS = {
    "open": "#4dabf7",
    "se-wip": "#ffc078",
    "l1-wip": "#9775fa",
    "l2": "#9775fa",
    "l3": "#51cf66",
    "l3-wip": "#51cf66",
    "tm to assign": "#ffa8cc",
    "need info": "#ffd43b",
    "need info from client": "#ffd43b",
    "client to validate": "#51cf66",
    "client to validate & close": "#51cf66",
    "resolved": "#868e96",
    "closed": "#868e96",
    "done": "#868e96",
    "completed": "#868e96",
}
SEVERITY_COLORS = {
    "blocker": "#ff6b6b",
    "critical": "#ff8c00",
    "high": "#ff6347",
    "medium": "#ffa94d",
    "low": "#69db7c",
    "info": "#868e96",
}
ENV_COLORS = {
    "STG": "#da77f2",
    "stg": "#da77f2",
    "staging": "#da77f2",
    "live": "#40c057",
    "production": "#40c057",
    "prod": "#40c057",
    "TBD": "#868e96",
}
ISSUE_COLORS = {
    "Bug": "#ff6b6b",
    "Feature": "#4dabf7",
    "Enhancement": "#69db7c",
    "Documentation": "#ffd43b",
    "Hotfix": "#ff8c00",
    "Other": "#868e96",
}
RESOLVED_KEYWORDS = ["resolved", "closed", "done", "completed", "client to validate & close"]

PRESETS = {
    "default": "[Sn]. [ID] | [Title] | [Issue Type] | [Status] | [Env] | [ETA]",
    "simple": "[Sn]. [ID] - [Title] - [Status]",
    "detailed": "[Sn]. [ID] | [Title] | [Severity] | [Status] | [Owner] | [Env] | [ETA] | [Issue Type]",
    "title_only": "[Sn]. [ID] | [Title]",
}


def load_db():
    if os.path.exists(DB_FILE):
        with open(DB_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}


def save_db(db):
    with open(DB_FILE, "w", encoding="utf-8") as f:
        json.dump(db, f, indent=2, ensure_ascii=False)


def load_gsheet_links():
    if os.path.exists(GSHEET_LINKS_FILE):
        with open(GSHEET_LINKS_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}


def save_gsheet_links(links):
    with open(GSHEET_LINKS_FILE, "w", encoding="utf-8") as f:
        json.dump(links, f, indent=2, ensure_ascii=False)


def read_csv(path):
    records = []
    with open(path, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            records.append(row)
    return records


def read_excel(path):
    try:
        import openpyxl
    except ImportError:
        import subprocess, sys
        subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
        import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue
        records.append(dict(zip(headers, row)))
    return records


def read_gsheet_public(url, sheet_name=""):
    url = url.strip()
    if "/edit" in url or "/form" in url:
        url = url.replace("/edit#gid=", "/export?format=csv&gid=")
        url = re.sub(r"/edit.*", "/export?format=csv", url)
        if "gid=" not in url:
            url += ("?" if "?" not in url else "&") + "format=csv"
    elif "export?format=csv" not in url:
        url = url.rstrip("/")
        url += "/export?format=csv"

    try:
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=30) as response:
            content = response.read().decode("utf-8-sig")
        from io import StringIO
        reader = csv.DictReader(StringIO(content))
        records = []
        for row in reader:
            if any(v for v in row.values()):
                records.append(row)
        return records, None
    except Exception as e:
        return [], str(e)


def extract_id(row):
    for key in ["ID", "#,Title,Project,ID", "#", "Ticket ID"]:
        val = row.get(key, "")
        if val:
            return str(val).lstrip("#").strip()
    for k, v in row.items():
        if v and isinstance(v, str) and re.match(r"#?\s*WOTASD-\d+", v.strip(), re.I):
            return str(v).lstrip("#").strip()
    return ""


def get_project_name(row):
    proj = row.get("Project", "")
    if proj:
        return str(proj).strip()
    title = (row.get("Title") or "").strip()
    match = re.match(r"([A-Z]{2,10}-\d+)", title)
    if match:
        prefix = match.group(1).split("-")[0]
        return prefix
    return "default"


def get_color(tag_type, value):
    colors = {
        "status": STATUS_COLORS,
        "severity": SEVERITY_COLORS,
        "env": ENV_COLORS,
        "issue": ISSUE_COLORS,
    }.get(tag_type, {})
    vl = (value or "").lower()
    for key, color in colors.items():
        if key in vl:
            return color
    return "#e9ecef"


def build_bullet(row, id_map, sn, fmt):
    raw_id = extract_id(row)
    rec = id_map.get(raw_id, {})
    mapping = {
        "Sn": sn, "#": sn,
        "ID": raw_id,
        "Title": (row.get("Title") or rec.get("Title") or "").strip().strip('"'),
        "Status": (row.get("Status") or rec.get("Status") or "").strip(),
        "Environment": (rec.get("Environment") or "").strip(),
        "Env": (rec.get("Environment") or "").strip(),
        "ETA": (rec.get("ETA") or "").strip(),
        "Issue Type": (rec.get("Issue Type") or "").strip(),
        "Severity": (row.get("Severity") or rec.get("Severity") or "").strip(),
        "Owner": (row.get("Owner") or rec.get("Owner") or "").strip(),
        "Project": (rec.get("Project") or row.get("Project") or "").strip(),
    }
    result = fmt
    for k, v in mapping.items():
        result = result.replace(f"[{k}]", str(v))
    return result


def generate_html_bullet(row, id_map, sn, fmt):
    raw_id = extract_id(row)
    rec = id_map.get(raw_id, {})
    status = (row.get("Status") or rec.get("Status") or "").strip()
    severity = (row.get("Severity") or rec.get("Severity") or "").strip().lower()
    env = (rec.get("Environment") or "").strip()
    issue_type = (rec.get("Issue Type") or "").strip()
    title = (row.get("Title") or rec.get("Title") or "").strip().strip('"')

    parts = fmt.split("|")
    html_parts = []
    for j, part_template in enumerate(parts):
        part = part_template.strip()
        rendered = part
        for k, v in {
            "Sn": sn, "#": sn, "ID": raw_id,
            "Title": title, "Status": status,
            "Environment": env, "Env": env, "ETA": (rec.get("ETA") or "").strip(),
            "Issue Type": issue_type, "Severity": (row.get("Severity") or rec.get("Severity") or "").strip(),
            "Owner": (row.get("Owner") or rec.get("Owner") or "").strip(),
            "Project": (rec.get("Project") or row.get("Project") or "").strip(),
        }.items():
            rendered = rendered.replace(f"[{k}]", str(v))

        color = "#e9ecef"
        if j == 1 and raw_id:
            color = "#da77f2"
        elif j == 2 and title:
            color = "#e9ecef"
        elif j == 3 and issue_type:
            color = get_color("issue", issue_type)
        elif j == 4 and status:
            color = get_color("status", status)
        elif j == 5 and env:
            color = get_color("env", env)
        elif j == 6 and env:
            color = get_color("env", env)
        elif j == 3 and not issue_type and status:
            color = get_color("status", status)
        elif j == 2 and not title and status:
            color = get_color("status", status)
        elif j == 2 and not title:
            color = "#e9ecef"

        html_parts.append(f'<span style="color:{color}">{rendered}</span>')

    return " | ".join(html_parts)


def generate_html_output(rows, id_map, fmt, title_text="Ticket Status Report"):
    ticket_rows = []
    for i, row in enumerate(rows, 1):
        raw_id = extract_id(row)
        rec = id_map.get(raw_id, {})
        status = (row.get("Status") or rec.get("Status") or "").strip()
        severity = (row.get("Severity") or rec.get("Severity") or "").strip().lower()
        env = (rec.get("Environment") or "").strip()
        issue_type = (rec.get("Issue Type") or "").strip()
        title = (row.get("Title") or rec.get("Title") or "").strip().strip('"')
        owner = (row.get("Owner") or rec.get("Owner") or "").strip()
        project = (rec.get("Project") or row.get("Project") or "").strip()

        s_color = get_color("status", status)
        sev_color = get_color("severity", severity)
        env_color = get_color("env", env)
        issue_color = get_color("issue", issue_type)

        line = generate_html_bullet(row, id_map, i, fmt)

        ticket_rows.append(f"""<div class="ticket-row">
  <span class="sn">{i}</span>
  <span class="id" style="color:#da77f2">{raw_id}</span>
  <span class="title">{title}</span>
  <span class="issue-type" style="color:{issue_color}">{issue_type}</span>
  <span class="status" style="color:{s_color}">{status}</span>
  <span class="env" style="color:{env_color}">{env}</span>
  <span class="eta" style="color:{env_color}">{rec.get('ETA','')}</span>
  <span class="severity" style="color:{sev_color}">{severity.title()}</span>
  <span class="owner">{owner}</span>
</div>""")

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>{title_text}</title>
<style>
* {{ box-sizing: border-box; margin: 0; padding: 0; }}
body {{
  background: #0a0a0a;
  color: #e9ecef;
  font-family: 'Segoe UI', Arial, sans-serif;
  padding: 40px;
}}
h1 {{
  color: #da77f2;
  font-size: 28px;
  margin-bottom: 8px;
  font-weight: 600;
  letter-spacing: -0.5px;
}}
.subtitle {{
  color: #868e96;
  font-size: 14px;
  margin-bottom: 30px;
}}
.ticket-row {{
  display: flex;
  gap: 0;
  padding: 8px 0;
  border-bottom: 1px solid #1a1a2e;
  font-size: 14px;
  align-items: center;
}}
.ticket-row:hover {{
  background: #111827;
}}
.ticket-row:first-of-type {{
  border-top: 2px solid #da77f2;
  border-bottom: 2px solid #da77f2;
  padding: 10px 0;
  font-weight: 600;
  color: #868e96;
  font-size: 12px;
  text-transform: uppercase;
  letter-spacing: 0.5px;
}}
.sn {{ width: 40px; color: #495057; flex-shrink: 0; }}
.id {{ width: 160px; flex-shrink: 0; font-family: Consolas, 'Courier New', monospace; font-size: 13px; }}
.title {{ flex: 2; min-width: 200px; color: #e9ecef; }}
.issue-type {{ width: 130px; flex-shrink: 0; font-weight: 600; font-size: 12px; text-transform: uppercase; letter-spacing: 0.5px; }}
.status {{ width: 180px; flex-shrink: 0; font-weight: 500; font-size: 13px; }}
.env {{ width: 60px; flex-shrink: 0; font-weight: 600; font-size: 12px; text-align: center; }}
.eta {{ width: 60px; flex-shrink: 0; font-size: 12px; text-align: center; }}
.severity {{ width: 90px; flex-shrink: 0; font-size: 12px; }}
.owner {{ width: 140px; flex-shrink: 0; color: #adb5bd; font-size: 13px; }}
.footer {{
  margin-top: 30px;
  color: #495057;
  font-size: 12px;
  text-align: center;
}}
</style>
</head>
<body>
<h1>{title_text}</h1>
<div class="subtitle">Generated: {__import__('datetime').datetime.now().strftime('%Y-%m-%d %H:%M')}</div>
<div class="ticket-row">
  <span class="sn">#</span>
  <span class="id">Ticket ID</span>
  <span class="title">Title</span>
  <span class="issue-type">Issue Type</span>
  <span class="status">Status</span>
  <span class="env">Env</span>
  <span class="eta">ETA</span>
  <span class="severity">Severity</span>
  <span class="owner">Owner</span>
</div>
{chr(10).join(ticket_rows)}
<div class="footer">Ticket Status Report &bull; Total: {len(rows)} tickets</div>
</body>
</html>"""


def get_projects(db):
    projects = {}
    for tid, data in db.items():
        if tid.startswith("_"):
            continue
        proj = (data.get("Project") or "default").strip()
        if not proj:
            proj = "default"
        proj_lower = proj.lower()
        if proj_lower not in projects:
            projects[proj_lower] = {
                "id": proj_lower,
                "label": proj,
                "tickets": [],
                "gsheet_url": "",
                "gsheet_enabled": False,
            }
        projects[proj_lower]["tickets"].append({**data, "tid": tid})
    gsheet_links = load_gsheet_links()
    for pid, link_data in gsheet_links.items():
        pl = pid.lower()
        if pl in projects:
            projects[pl]["gsheet_url"] = link_data.get("url", "")
            projects[pl]["gsheet_enabled"] = link_data.get("enabled", False)
        elif link_data.get("enabled"):
            projects[pl] = {
                "id": pl,
                "label": pid,
                "tickets": [],
                "gsheet_url": link_data.get("url", ""),
                "gsheet_enabled": True,
            }
    return list(projects.values())


def import_records_to_db(records, project, batch_env):
    db = load_db()
    new_count = 0
    updated_count = 0
    current_ids = set()
    missing_ids = []

    for row in records:
        raw_id = extract_id(row)
        if not raw_id:
            continue
        current_ids.add(raw_id)
        rec = {**row, "imported": True, "Project": project}
        if batch_env:
            rec["Environment"] = batch_env
            rec["ETA"] = batch_env

        if raw_id in db:
            db[raw_id] = {**db[raw_id], **rec}
            updated_count += 1
        else:
            db[raw_id] = rec
            new_count += 1

    for tid, data in db.items():
        if tid.startswith("_"):
            continue
        if data.get("Project", "").lower() != project.lower():
            continue
        if tid not in current_ids:
            missing_ids.append({**data, "tid": tid})

    save_db(db)
    return new_count, updated_count, current_ids, missing_ids


def import_file_to_db(file_path, project, batch_env):
    ext = Path(file_path).suffix.lower()
    if ext in [".xlsx", ".xls"]:
        records = read_excel(file_path)
    elif ext == ".csv":
        records = read_csv(file_path)
    else:
        return 0, 0, [], f"Unsupported: {ext}"
    new_c, updated_c, _, missing = import_records_to_db(records, project, batch_env)
    return new_c, updated_c, missing, None


def import_gsheet_to_db(url, project, batch_env):
    records, err = read_gsheet_public(url)
    if err:
        return 0, 0, [], err
    if not records:
        return 0, 0, [], "No records found in sheet"
    new_c, updated_c, _, missing = import_records_to_db(records, project, batch_env)
    return new_c, updated_c, missing, None


@app.route("/")
def index():
    db = load_db()
    gsheet_links = load_gsheet_links()
    presets = list(PRESETS.keys())
    current_fmt = db.get("_settings", {}).get("ppt_format", "default")
    projects = get_projects(db)
    return render_template(
        "index.html",
        db=db,
        projects=projects,
        presets=presets,
        current_fmt=current_fmt,
        gsheet_links=gsheet_links,
        STATUS_COLORS=json.dumps(STATUS_COLORS),
        SEVERITY_COLORS=json.dumps(SEVERITY_COLORS),
        ENV_COLORS=json.dumps(ENV_COLORS),
        ISSUE_COLORS=json.dumps(ISSUE_COLORS),
    )


@app.route("/api/import", methods=["POST"])
def api_import():
    project = request.form.get("project", "default").strip()
    batch_env = request.form.get("env", "STG").strip()
    file_path = None

    if "file" in request.files:
        f = request.files["file"]
        if f.filename:
            file_path = os.path.join("/tmp", secure_filename(f.filename))
            f.save(file_path)
    elif request.form.get("path"):
        file_path = request.form["path"]

    if not file_path or not os.path.exists(file_path):
        return jsonify({"success": False, "error": f"File not found: {file_path}"}), 400

    new_c, updated_c, missing, err = import_file_to_db(file_path, project, batch_env)
    if err:
        return jsonify({"success": False, "error": err}), 400

    return jsonify({
        "success": True,
        "imported": new_c,
        "updated": updated_c,
        "missing": missing,
        "missing_count": len(missing),
        "total": len(load_db()),
    })


@app.route("/api/gsheet/import", methods=["POST"])
def api_gsheet_import():
    data = request.get_json()
    url = data.get("url", "").strip()
    project = data.get("project", "default").strip()
    batch_env = data.get("env", "STG").strip()

    if not url:
        return jsonify({"success": False, "error": "URL required"}), 400

    new_c, updated_c, missing, err = import_gsheet_to_db(url, project, batch_env)
    if err:
        return jsonify({"success": False, "error": err}), 400

    return jsonify({
        "success": True,
        "imported": new_c,
        "updated": updated_c,
        "missing": missing,
        "missing_count": len(missing),
        "total": len(load_db()),
    })


@app.route("/api/gsheet/links", methods=["GET"])
def api_gsheet_links_get():
    return jsonify(load_gsheet_links())


@app.route("/api/gsheet/links", methods=["POST"])
def api_gsheet_links_save():
    links = request.get_json() or {}
    save_gsheet_links(links)
    return jsonify({"success": True, "saved": True})


@app.route("/api/tickets", methods=["GET"])
def api_tickets():
    return jsonify(load_db())


@app.route("/api/ticket/<ticket_id>", methods=["GET"])
def api_ticket_get(ticket_id):
    db = load_db()
    return jsonify(db.get(ticket_id, {}))


@app.route("/api/ticket/<ticket_id>", methods=["POST"])
def api_ticket_update(ticket_id):
    db = load_db()
    if ticket_id not in db:
        return jsonify({"success": False, "error": "Not found"}), 404
    data = request.get_json()
    db[ticket_id].update(data)
    save_db(db)
    return jsonify({"success": True, "ticket": db[ticket_id]})


@app.route("/api/ticket/<ticket_id>", methods=["DELETE"])
def api_ticket_delete(ticket_id):
    db = load_db()
    if ticket_id in db:
        del db[ticket_id]
        save_db(db)
    return jsonify({"success": True})


@app.route("/api/project/<project_id>", methods=["DELETE"])
def api_project_delete(project_id):
    db = load_db()
    proj = project_id.lower()
    deleted = 0
    for tid in list(db.keys()):
        if tid.startswith("_"):
            continue
        if db[tid].get("Project", "").lower() == proj:
            del db[tid]
            deleted += 1
    save_db(db)
    return jsonify({"success": True, "deleted": deleted})


@app.route("/api/bulk/resolve", methods=["POST"])
def api_bulk_resolve():
    data = request.get_json()
    actions = data.get("actions", {})
    db = load_db()
    counts = {"resolved": 0, "skipped": 0, "removed": 0}
    for tid, action in actions.items():
        if tid not in db:
            continue
        if action == "resolve":
            db[tid]["Status"] = "Resolved"
            counts["resolved"] += 1
        elif action == "remove":
            del db[tid]
            counts["removed"] += 1
        elif action == "skip":
            counts["skipped"] += 1
    save_db(db)
    return jsonify({"success": True, **counts})


@app.route("/api/generate/ppt", methods=["POST"])
def api_generate_ppt():
    data = request.get_json()
    fmt_name = data.get("format", "default")
    if fmt_name == "custom":
        fmt = data.get("custom_format", PRESETS["default"])
    else:
        fmt = PRESETS.get(fmt_name, PRESETS["default"])

    tickets_data = data.get("tickets", [])
    project_filter = data.get("project", "all").strip().lower()
    title_text = data.get("title", "Ticket Status Report")
    include_html = data.get("html", True)
    include_plain = data.get("plain", True)

    db = load_db()
    id_map = {k.lstrip("#").strip(): v for k, v in db.items()}
    id_map.update(db)

    filtered = tickets_data
    if project_filter and project_filter != "all":
        filtered = [t for t in tickets_data if (t.get("Project") or "default").lower() == project_filter]

    plain_lines = [build_bullet(row, id_map, i + 1, fmt) for i, row in enumerate(filtered)]

    result = {"success": True, "lines": len(plain_lines)}

    if include_plain:
        with open(PPT_PLAIN, "w", encoding="utf-8") as f:
            f.write("\n".join(plain_lines))
        result["plain_file"] = PPT_PLAIN

    if include_html:
        html = generate_html_output(filtered, id_map, fmt, title_text)
        with open(PPT_HTML, "w", encoding="utf-8") as f:
            f.write(html)
        result["html_file"] = PPT_HTML

        html_all = generate_html_output(
            [row for row in db.values() if not str(row).startswith("_")],
            id_map, fmt, title_text
        )
        with open(PPT_HTML_ALL, "w", encoding="utf-8") as f:
            f.write(html_all)
        result["html_all_file"] = PPT_HTML_ALL

    return jsonify(result)


@app.route("/api/download/<filetype>", methods=["GET"])
def api_download(filetype):
    path_map = {
        "plain": PPT_PLAIN,
        "html": PPT_HTML,
        "all": PPT_HTML_ALL,
    }
    path = path_map.get(filetype)
    if not path or not os.path.exists(path):
        return jsonify({"success": False, "error": "File not found. Generate first."}), 404
    name_map = {
        "plain": "ppt_output.txt",
        "html": "ppt_output.html",
        "all": "ppt_output_all.html",
    }
    return send_file(path, as_attachment=True, download_name=name_map.get(filetype, "output.txt"))


@app.route("/api/preview/ppt", methods=["POST"])
def api_preview_ppt():
    data = request.get_json()
    fmt_name = data.get("format", "default")
    if fmt_name == "custom":
        fmt = data.get("custom_format", PRESETS["default"])
    else:
        fmt = PRESETS.get(fmt_name, PRESETS["default"])

    tickets_data = data.get("tickets", [])
    project_filter = data.get("project", "all").strip().lower()
    limit = data.get("limit", 10)

    db = load_db()
    id_map = {k.lstrip("#").strip(): v for k, v in db.items()}
    id_map.update(db)

    filtered = tickets_data
    if project_filter and project_filter != "all":
        filtered = [t for t in tickets_data if (t.get("Project") or "default").lower() == project_filter]

    lines = [build_bullet(row, id_map, i + 1, fmt) for i, row in enumerate(filtered[:limit])]
    more = len(filtered) - limit
    preview = "\n".join(lines)
    if more > 0:
        preview += f"\n... and {more} more"
    return jsonify({"success": True, "preview": preview, "total": len(filtered)})


@app.route("/api/project/create", methods=["POST"])
def api_project_create():
    data = request.get_json()
    name = data.get("name", "").strip()
    if not name:
        return jsonify({"success": False, "error": "Project name required"}), 400
    return jsonify({"success": True, "project": name.lower(), "label": name})


@app.route("/api/settings", methods=["GET"])
def api_settings():
    db = load_db()
    return jsonify(db.get("_settings", {}))


@app.route("/api/settings", methods=["POST"])
def api_settings_save():
    db = load_db()
    data = request.get_json()
    db["_settings"] = {**(db.get("_settings", {})), **data}
    save_db(db)
    return jsonify({"success": True})


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)